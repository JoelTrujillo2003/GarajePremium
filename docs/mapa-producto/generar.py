"""Genera el mapa editable y una copia de la presentación con el mapa como anexo.
Uso: python3 docs/mapa-producto/generar.py /ruta/a/presentacion.pptx
"""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT=Path(__file__).resolve().parent
ORIGINAL=Path(sys.argv[1]) if len(sys.argv)>1 else None
NAVY=RGBColor(38,52,68)
NAVY_DARK=RGBColor(22,34,47)
SLATE=RGBColor(112,133,153)
PALE=RGBColor(237,241,245)
BLUE=RGBColor(223,232,242)
BLUE_DEEP=RGBColor(207,219,231)
ORANGE=RGBColor(233,162,88)
WHITE=RGBColor(255,255,255)


def text(slide,x,y,w,h,s,size=13,bold=False,color=NAVY,align=None):
    shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.01);tf.margin_right=Inches(.01)
    tf.margin_top=Inches(.01);tf.margin_bottom=Inches(.01)
    p=tf.paragraphs[0];p.text=s;p.space_before=Pt(0);p.space_after=Pt(0)
    if align:p.alignment=align
    for run in p.runs:
        run.font.name='Aptos';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color
    return shape


def card(slide,x,y,w,h,title,body,tag=None,fill=PALE):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background()
    if tag:text(slide,x+.16,y+.14,w-.3,.24,tag,10,True,SLATE)
    text(slide,x+.16,y+(.49 if tag else .26),w-.32,.38,title,17,True,NAVY)
    text(slide,x+.16,y+(.93 if tag else .75),w-.32,h-(1.02 if tag else .86),body,12,False,NAVY)
    return sh


def arrow(slide,x1,y1,x2,y2,color=SLATE):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=color;sh.line.width=Pt(2.0)
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    # Chevron arrowhead as its own editable shape, with no SVG dependency.
    tip=slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x2-.11), Inches(y2-.08), Inches(.13), Inches(.16))
    tip.fill.solid();tip.fill.fore_color.rgb=color;tip.line.fill.background()


def add_map(prs):
    slide=prs.slides.add_slide(prs.slide_layouts[min(6, len(prs.slide_layouts)-1)])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=WHITE
    text(slide,.59,.40,12.2,.52,'Mapa del producto',30,True,NAVY)
    text(slide,.61,.97,12,.30,'Del catálogo público a la guía de mantención',14,False,SLATE)
    text(slide,.61,1.44,12,.22,'RECORRIDO DE LA PERSONA',10,True,SLATE)
    xs=[.61,3.72,6.83,9.94]
    cards=[
        ('01  PÚBLICO','Inicio','Valor del producto,\nplanes y acceso.'),
        ('02  PÚBLICO','Catálogo','Auto o moto; marca →\nmodelo → año → km.'),
        ('03  IDENTIDAD','Cuenta','Registro o sesión\ncon Supabase Auth.'),
        ('04  PRIVADO','Mi garage','Fichas asociadas\na cada usuario.')
    ]
    for i,(tag,title,body) in enumerate(cards):
        card(slide,xs[i],1.76,2.78,1.65,title,body,tag,BLUE if i in (1,3) else PALE)
        if i<3:arrow(slide,xs[i]+2.81,2.58,xs[i+1]-.09,2.58)
    text(slide,.61,3.70,12,.22,'CAPACIDADES Y CAMINOS',10,True,SLATE)
    card(slide,.61,4.02,2.78,1.46,'Free','Hasta 2 vehículos.\nConsulta tu garage.','PLAN DE ENTRADA',PALE)
    card(slide,3.72,4.02,2.78,1.46,'Premium','Garage sin límite.\nAcceso a la guía IA.','PLAN AMPLIADO',BLUE)
    card(slide,6.83,4.02,2.78,1.46,'Orientación IA','Vehículo + observaciones\n→ reporte con Gemini.','SOLO PREMIUM',PALE)
    card(slide,9.94,4.02,2.78,1.46,'Hablemos','Consultas y sugerencias\nregistradas en Supabase.','CANAL PÚBLICO',BLUE)
    arrow(slide,6.52,4.75,6.78,4.75,ORANGE)
    text(slide,.62,5.60,12,.23,'CÓMO SE IMPLEMENTA',10,True,SLATE)
    band=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.61),Inches(5.91),Inches(12.11),Inches(.87))
    band.fill.solid();band.fill.fore_color.rgb=NAVY_DARK;band.line.fill.background()
    parts=[('NAVEGADOR','HTML · CSS · JavaScript'),('SERVIDOR EN RENDER','Node.js · Express'),('DATOS Y ACCESO','Supabase · PostgreSQL · RLS'),('GUÍA','Gemini desde el servidor')]
    for i,(head,sub) in enumerate(parts):
        x=.81+i*3.02
        text(slide,x,6.02,2.82,.22,head,10,True,WHITE)
        text(slide,x,6.29,2.82,.30,sub,11,False,RGBColor(214,225,236))
        if i<3:
            text(slide,x+2.82,6.22,.20,.22,'→',11,True,ORANGE)
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.61),Inches(7.03),Inches(12.11),Inches(.025))
    line.fill.solid();line.fill.fore_color.rgb=ORANGE;line.line.fill.background()
    text(slide,.61,7.14,6,.16,'GARAGE PREMIUM MODERNO   /   HITO 1',8,True,SLATE)
    text(slide,11.80,7.14,.93,.16,'ANEXO / MAPA',8,True,SLATE,PP_ALIGN.RIGHT)
    return slide

standalone=Presentation();standalone.slide_width=Inches(13.333333);standalone.slide_height=Inches(7.5)
add_map(standalone);standalone.save(OUT/'Mapa-de-producto-editable.pptx')
if ORIGINAL:
    combined=Presentation(str(ORIGINAL));add_map(combined)
    combined.save(OUT/'Presentacion-con-mapa-de-producto.pptx')
print('Mapa editable generado en',OUT)
